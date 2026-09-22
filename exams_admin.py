# -*- coding: utf-8 -*-
"""教师考试管理后台（挂 /admin/exams，复用统一后台的密码 session）。

功能：
- 新建/删除考试，动态录入选择题；
- 设置练习端口开放窗口（按班、按时段）；
- 查看单场考试提交情况、缺考名单、待复核答卷；
- 导出单场成绩表；导出"课前 vs 课后"配对对比表（含提升值、随机抽查名单）。
"""
import os
from functools import wraps
from flask import (Blueprint, request, jsonify, render_template, send_file)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import exams
from students import list_classes

bp = Blueprint('exams_admin', __name__)

OUT_DIR = os.path.join(exams.BASE, '成绩导出')
os.makedirs(OUT_DIR, exist_ok=True)

HDR_FILL = PatternFill('solid', fgColor='1B3A5C')
HDR_FONT = Font(color='FFFFFF', bold=True, size=11)
WARN_FILL = PatternFill('solid', fgColor='FDEBD0')
THIN = Border(*[Side(style='thin')] * 4)
CENTER = Alignment(horizontal='center', vertical='center')


def admin_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        from flask import session, redirect
        if not session.get('admin_ok'):
            if request.path.startswith('/admin/exams/api'):
                return jsonify({'error': '请先登录后台'}), 401
            return redirect('/admin/login')
        return fn(*args, **kwargs)
    return wrapper


# ───────── 页面与考试增删 ─────────
@bp.route('/admin/exams')
@admin_required
def page():
    embedded = request.args.get('embedded') == '1'
    return render_template('admin_exams.html', classes=list_classes(), embedded=embedded)


@bp.route('/admin/exams/api/list')
@admin_required
def api_list():
    out = []
    for e in exams.list_exams():
        subs = exams.list_submissions(e['id'])
        expected = len(exams.expected_students(e))
        graded = [s for s in subs if s.get('status') == 'graded']
        out.append({'id': e['id'], 'name': e['name'], 'module': e['module'],
                    'classes': e['classes'], 'start': e['start'], 'end': e['end'],
                    'quiz_count': len(e.get('quiz') or []),
                    'submitted': len(subs), 'expected': expected,
                    'graded': len(graded),
                    'need_review': len(subs) - len(graded)})
    return jsonify({'exams': out})


@bp.route('/admin/exams/api/create', methods=['POST'])
@admin_required
def api_create():
    d = request.get_json(force=True)

    # 选择题：前端按 q0/q1 提交
    quiz = []
    n = int(d.get('quiz_count') or 0)
    for i in range(n):
        q = (d.get(f'q{i}_t') or '').strip()
        opts = [d.get(f'q{i}_o{j}', '') for j in range(4)]
        ans = d.get(f'q{i}_a')
        if not q or any(not o.strip() for o in opts) or ans not in ('0', '1', '2', '3'):
            return jsonify({'error': f'第{i+1}题信息不完整（题干/4个选项/正确答案）'}), 400
        quiz.append({'q': q, 'options': opts, 'answer': int(ans)})

    classes = d.get('classes') or []
    exam, err = exams.create_exam(
        d.get('name'), d.get('module'), classes,
        d.get('start'), d.get('end'), quiz=quiz,
        quiz_max=int(d.get('quiz_max') or 20), note=d.get('note'))
    if err:
        return jsonify({'error': err}), 400
    return jsonify({'ok': True, 'id': exam['id'], 'link': f'/exam/{exam["id"]}'})


@bp.route('/admin/exams/api/delete', methods=['POST'])
@admin_required
def api_delete():
    eid = (request.get_json(force=True) or {}).get('id')
    err = exams.delete_exam(eid)
    return jsonify({'ok': not err, 'error': err})


# ───────── 练习端口窗口 ─────────
@bp.route('/admin/exams/api/practice', methods=['POST'])
@admin_required
def api_practice_set():
    d = request.get_json(force=True)
    err = exams.set_practice_window(d.get('module'), d.get('classes'),
                                    d.get('start'), d.get('end'))
    return jsonify({'ok': not err, 'error': err})


@bp.route('/admin/exams/api/practice/clear', methods=['POST'])
@admin_required
def api_practice_clear():
    d = request.get_json(force=True) or {}
    exams.clear_practice_window(d.get('module', 'xray'))
    return jsonify({'ok': True})


@bp.route('/admin/exams/api/practice')
@admin_required
def api_practice_get():
    return jsonify({'window': exams.get_practice_window('xray')})


# ───────── 单场明细 ─────────
@bp.route('/admin/exams/api/detail/<exam_id>')
@admin_required
def api_detail(exam_id):
    exam = exams.get_exam(exam_id)
    if not exam:
        return jsonify({'error': '考试不存在'}), 404
    subs = exams.list_submissions(exam_id, include_teacher=False)
    submitted_sids = {s['sid'] for s in subs}
    absent = [{'sid': st['sid'], 'name': st['name'], 'class': st['class']}
              for st in exams.expected_students(exam)
              if st['sid'] not in submitted_sids]
    return jsonify({'exam': {k: exam.get(k) for k in
                             ('id', 'name', 'start', 'end', 'classes', 'quiz_max', 'skill_max')},
                    'submissions': subs, 'absent': absent})


# ───────── 导出：单场成绩 ─────────
def _style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(1, c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = CENTER


def _save_wb(wb, filename):
    path = os.path.join(OUT_DIR, filename)
    wb.save(path)
    return path


@bp.route('/admin/exams/api/export/<exam_id>')
@admin_required
def api_export(exam_id):
    exam = exams.get_exam(exam_id)
    if not exam:
        return jsonify({'error': '考试不存在'}), 404
    subs = {s['sid']: s for s in exams.list_submissions(exam_id)}

    wb = Workbook(); ws = wb.active; ws.title = '成绩'
    cols = ['学号', '姓名', '班级']
    if exam.get('quiz'):
        cols += ['选择题', 'X光片判读']
    cols += ['总分', '状态', '提交时间']
    ws.append(cols)

    for st in exams.expected_students(exam):
        s = subs.get(st['sid'])
        row = [st['sid'], st['name'], st['class']]
        if s:
            if exam.get('quiz'):
                row += [s.get('quiz_score'), s.get('skill_score')]
            row += [s.get('total_score'),
                    '已评分' if s.get('status') == 'graded' else '待复核',
                    s.get('submitted_at')]
            if s.get('status') != 'graded':
                row_n = ws.max_row + 1
        else:
            if exam.get('quiz'):
                row += ['', '']
            row += ['', '缺考', '']
        ws.append(row)
        if s and s.get('status') != 'graded':
            for c in range(1, len(cols) + 1):
                ws.cell(ws.max_row, c).fill = WARN_FILL

    _style_header(ws, len(cols))
    for col, w in zip('ABCDEFGHI', [20, 12, 22, 9, 11, 8, 9, 20]):
        ws.column_dimensions[col].width = w
    path = _save_wb(wb, f'{exam["name"]}_成绩表.xlsx')
    return send_file(path, as_attachment=True)


# ───────── 导出：课前 vs 课后配对对比 ─────────
@bp.route('/admin/exams/api/compare')
@admin_required
def api_compare():
    pre_id, post_id = request.args.get('pre'), request.args.get('post')
    pre_exam, post_exam = exams.get_exam(pre_id), exams.get_exam(post_id)
    if not pre_exam or not post_exam:
        return jsonify({'error': '请先选择课前和课后两场考试'}), 400

    pre = {s['sid']: s for s in exams.list_submissions(pre_id)}
    post = {s['sid']: s for s in exams.list_submissions(post_id)}
    sids = sorted(set(pre) | set(post))

    wb = Workbook(); ws = wb.active; ws.title = '课前课后对比'
    has_quiz = bool(pre_exam.get('quiz') or post_exam.get('quiz'))
    head = ['学号', '姓名', '班级']
    if has_quiz:
        head += ['课前选择', '课后选择']
    head += ['课前判读', '课后判读', '课前总分', '课后总分', '提升值', '备注']
    ws.append(head)

    gains = []
    for sid in sids:
        a, b = pre.get(sid), post.get(sid)
        ref = b or a
        row = [sid, ref.get('name', ''), ref.get('class', '')]
        note = []
        if has_quiz:
            row += [a.get('quiz_score') if a else '',
                    b.get('quiz_score') if b else '']
        row += [a.get('skill_score') if a else '',
                b.get('skill_score') if b else '',
                a.get('total_score') if a else '',
                b.get('total_score') if b else '']
        ta, tb = (a or {}).get('total_score'), (b or {}).get('total_score')
        if isinstance(ta, (int, float)) and isinstance(tb, (int, float)):
            row.append(round(tb - ta, 1)); gains.append(tb - ta)
        else:
            row.append('')
        if not a: note.append('缺课前')
        if not b: note.append('缺课后')
        if (a and a.get('status') != 'graded') or (b and b.get('status') != 'graded'):
            note.append('有待复核')
        row.append('、'.join(note))
        ws.append(row)
        if note:
            for c in range(1, len(head) + 1):
                ws.cell(ws.max_row, c).fill = WARN_FILL

    _style_header(ws, len(head))
    widths = [20, 12, 22, 9, 9, 10, 10, 9, 9, 8, 14]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = 'A2'

    # 汇总行
    n_both = len(gains)
    avg_gain = round(sum(gains) / n_both, 2) if n_both else 0
    ws.append([])
    ws.append(['配对人数', n_both, '', '', '', '', '', '', '平均提升', avg_gain])

    # Sheet2：随机抽查名单（用于人工盲评验AI）
    ws2 = wb.create_sheet('抽查名单')
    ws2.append(['序号', '学号', '姓名', '班级', '课前总分', '课后总分'])
    # 伪随机：按学号均匀抽样，无需随机数种子，结果可复现
    paired = sorted([sid for sid in sids if sid in pre and sid in post
                     and isinstance(pre[sid].get('total_score'), (int, float))])
    k = min(25, len(paired))
    picked = [paired[round(i * (len(paired) - 1) / max(1, k - 1))]
              for i in range(k)] if k else []
    seen2 = set(); picked2 = []
    for sid in picked:                       # 去重（等距抽样可能重复）
        if sid not in seen2:
            seen2.add(sid); picked2.append(sid)
    for i, sid in enumerate(picked2, 1):
        ws2.append([i, sid, pre[sid].get('name'), pre[sid].get('class'),
                    pre[sid].get('total_score'), post[sid].get('total_score')])
    _style_header(ws2, 6)
    for col, w in zip('ABCDEF', [6, 20, 12, 22, 10, 10]):
        ws2.column_dimensions[col].width = w

    path = _save_wb(wb, f'{pre_exam["name"]}_课前课后对比.xlsx')
    return send_file(path, as_attachment=True)
