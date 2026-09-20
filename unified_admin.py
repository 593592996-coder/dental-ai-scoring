# -*- coding: utf-8 -*-
"""统一教师后台：全板块成绩总表 + 学生画像 + 可勾选导出向导。

- 密码保护 /admin、/dashboard、/api/export、/train/grades；
- 数据来自 gradebook（5 个操作评分 + 临床思维，按学号并表）；
- 导出两种：学校格式 .xls（板块→学校列可勾选映射，综合=所填板块均分）、全量总表 .xlsx。
"""
import os
from datetime import datetime
from functools import wraps
from flask import (Blueprint, render_template, request, jsonify, session,
                   redirect, send_file, abort)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import gradebook
from gradebook import MODULES, MODULE_NAME, SCORING_MODULES, collect, class_rows, composite
import admin_data
from students import list_classes, class_template_map, fill_workbook, SCORE_COLS

bp = Blueprint('unified_admin', __name__)

BASE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(BASE, '成绩导出')
os.makedirs(OUT_DIR, exist_ok=True)

_PREFIX_PATH = {'class2': '', 'endo': 'endo_', 'xray': 'xray_',
                'crown_ant': 'crown_ant_', 'crown_post': 'crown_post_',
                'impression': 'impression_'}


def admin_password():
    pwd = os.environ.get('ADMIN_PASSWORD')
    if pwd:
        return pwd.strip()
    pwf = os.path.join(BASE, 'admin_password.txt')
    if os.path.exists(pwf):
        try:
            return open(pwf, encoding='utf-8').read().strip()
        except Exception:
            pass
    return 'aimei2026'


# 同时保护老看板与老导出，避免学生绕过
_PROTECT_PREFIX = ('/admin', '/dashboard', '/train/grades')
_PROTECT_EXACT = ('/api/export',)


@bp.before_app_request
def _guard():
    p = request.path
    protected = p.startswith(_PROTECT_PREFIX) or p in _PROTECT_EXACT
    if not protected or p.startswith('/admin/login'):
        return None
    if session.get('admin_ok'):
        return None
    if p.startswith('/admin/api') or p.startswith('/api/'):
        return jsonify({'error': '请先在 /admin/login 输入后台密码'}), 401
    return redirect('/admin/login')


@bp.route('/admin/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        if (request.form.get('password') or '').strip() == admin_password():
            session['admin_ok'] = True
            return redirect('/admin')
        return render_template('admin_login.html', error='密码不正确，请重试')
    if session.get('admin_ok'):
        return redirect('/admin')
    return render_template('admin_login.html', error='')


@bp.route('/admin/logout', methods=['GET', 'POST'])
def logout():
    session.pop('admin_ok', None)
    return redirect('/admin/login')


@bp.route('/admin')
def home():
    return render_template('admin_unified.html')


# ───────── 数据 API ─────────
@bp.route('/admin/api/overview')
def overview():
    klass = request.args.get('class', '').strip()
    rows, unmatched, junk = class_rows(klass)
    submitted = [r for r in rows if r['done_modules']]

    def avg(key):
        vs = [r[key] for r in submitted if r.get(key) is not None]
        return round(sum(vs) / len(vs), 1) if vs else None

    # 默认综合：六个板块简单平均
    for r in rows:
        c, have, miss = composite(r, [k for k, _ in MODULES])
        r['composite'] = c

    dist = {'优(≥85)': 0, '中(60-84)': 0, '待提高(<60)': 0, '未完成': 0}
    for r in rows:
        v = r['composite']
        if v is None:
            dist['未完成'] += 1
        elif v >= 85:
            dist['优(≥85)'] += 1
        elif v >= 60:
            dist['中(60-84)'] += 1
        else:
            dist['待提高(<60)'] += 1

    return jsonify({
        'classes': list_classes(), 'class': klass,
        'modules': [{'key': k, 'name': n} for k, n in MODULES],
        'total': len(rows), 'submitted': len(submitted),
        'avg': {k: avg(k) for k, _ in MODULES},
        'dist': dist, 'rows': rows,
        'unmatched': unmatched, 'junk': junk,
    })


@bp.route('/admin/api/student/<sid>')
def student(sid):
    rows_all, _, _ = class_rows('')
    rec = next((r for r in rows_all if r['sid'] == sid), None)
    if not rec:
        return jsonify({'error': '无此学生'}), 404
    # 操作评分报告链接
    links = {}
    for mod, rid in (rec.get('report_ids') or {}).items():
        if rid:
            links[MODULE_NAME[mod]] = f'/report/{_PREFIX_PATH.get(mod, "")}{rid}'
    out = {'student': {k: rec[k] for k in ('sid', 'name', 'class')},
           'modules': {k: rec.get(k) for k, _ in MODULES},
           'report_links': links,
           'thinking_detail': rec.get('thinking_detail')}
    # 思维训练原始作答
    try:
        td = admin_data.student_detail(sid)
        if td:
            out['thinking'] = {'summary': td['summary'], 'events': td['events'],
                               'dim_meta': td['dim_meta']}
    except Exception:
        pass
    return jsonify(out)


# ───────── 导出 ─────────
def _class_template(klass):
    tmap = class_template_map()
    if klass and klass in tmap:
        return tmap[klass], klass
    if len(tmap) == 1:
        k = next(iter(tmap))
        return tmap[k], k
    return None, klass


@bp.route('/admin/api/export_school', methods=['POST'])
def export_school():
    """学校格式：{class, mapping:{学校列名: 板块key}}；综合成绩=被映射板块的简单平均。"""
    d = request.get_json() or {}
    klass = (d.get('class') or '').strip()
    mapping = d.get('mapping') or {}
    # 只保留合法列与合法板块
    col2mod = {c: m for c, m in mapping.items() if c in SCORE_COLS and m in dict(MODULES)}
    if '综合成绩' in col2mod:
        col2mod.pop('综合成绩')  # 综合由系统算，不直接绑定单一板块
    if not col2mod:
        return jsonify({'error': '请至少勾选一个板块对应到学校成绩列'}), 400

    src, klass = _class_template(klass)
    if not src:
        return jsonify({'error': '找不到该班学校模板，请把教务空表放入 花名册/'}), 400

    rows, _, _ = class_rows(klass)
    used_mods = list(dict.fromkeys(col2mod.values()))
    score_map = {}
    for r in rows:
        comp, have, _ = composite(r, used_mods)
        one = {'综合成绩': comp} if comp is not None else {}
        for col, mod in col2mod.items():
            if r.get(mod) is not None:
                one[col] = r[mod]
        if one:
            score_map[r['sid']] = one
    stamp = datetime.now().strftime('%Y%m%d_%H%M')
    safe = klass.replace('/', '_').replace(' ', '') or '全部'
    dst = os.path.join(OUT_DIR, f'学校成绩表_{safe}_{stamp}.xls')
    written, extra = fill_workbook(src, dst, score_map)
    return jsonify({'ok': True, 'written': written, 'extra': extra,
                    'file': os.path.basename(dst), 'download': f'/admin/download?f={os.path.basename(dst)}'})


@bp.route('/admin/api/export_full', methods=['POST'])
def export_full():
    """全量总表 xlsx：{class, include:[板块key...]}；含每板块分 + 综合(所选均分)。"""
    d = request.get_json() or {}
    klass = (d.get('class') or '').strip()
    include = [m for m in (d.get('include') or []) if m in dict(MODULES)]
    if not include:
        return jsonify({'error': '请勾选至少一个板块'}), 400
    rows, unmatched, _ = class_rows(klass)

    wb = Workbook(); ws = wb.active; ws.title = '成绩总表'
    hdr = ['学号', '姓名', '班级'] + [MODULE_NAME[k] for k in include] + [f'综合成绩（{len(include)}板块均分）', '已完成板块数']
    ws.append(hdr)
    hf = Font(color='FFFFFF', bold=True); fill = PatternFill('solid', fgColor='1B3A5C')
    thin = Side(style='thin', color='BBBBBB'); border = Border(thin, thin, thin, thin)
    for c in ws[1]:
        c.font = hf; c.fill = fill; c.alignment = Alignment(horizontal='center'); c.border = border
    for r in rows:
        comp, have, _ = composite(r, include)
        ws.append([r['sid'], r['name'], r.get('class', '')] +
                  [r.get(k) if r.get(k) is not None else '' for k in include] +
                  [comp if comp is not None else '', have])
    # 未匹配（无学号历史成绩）单列一页
    ws2 = wb.create_sheet('待核对(无学号历史)')
    ws2.append(['姓名', '班级', '板块', '分数', '报告ID', '时间'])
    for u in unmatched:
        if klass and u.get('class') != klass:
            continue
        ws2.append([u['name'], u['class'], u['module_name'], u['score'], u['report_id'], u['ts']])
    for i, w in enumerate([18, 12, 22] + [12] * len(include) + [18, 10], 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else 'A' + chr(64 + i - 26)].width = w
    ws.freeze_panes = 'A2'

    stamp = datetime.now().strftime('%Y%m%d_%H%M')
    safe = klass.replace('/', '_').replace(' ', '') or '全部'
    fn = f'成绩总表_{safe}_{stamp}.xlsx'
    wb.save(os.path.join(OUT_DIR, fn))
    return jsonify({'ok': True, 'file': fn, 'download': f'/admin/download?f={fn}',
                    'unmatched': len(unmatched)})


@bp.route('/admin/download')
def download():
    f = os.path.basename(request.args.get('f', ''))
    path = os.path.join(OUT_DIR, f)
    if not f or not os.path.exists(path):
        abort(404)
    return send_file(path, as_attachment=True)
